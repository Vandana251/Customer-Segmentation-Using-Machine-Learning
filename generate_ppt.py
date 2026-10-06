import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # Slate 900
    LIGHT_BG = RGBColor(248, 250, 252)  # Slate 50
    CARD_BG = RGBColor(255, 255, 255)   # White
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    DARK_CARD = RGBColor(30, 41, 59)    # Slate 800
    
    NAVY = RGBColor(15, 44, 89)
    CYAN = RGBColor(14, 165, 233)       # Sky 500
    TEAL = RGBColor(20, 184, 166)       # Teal 500
    CORAL = RGBColor(244, 63, 94)       # Rose 500
    AMBER = RGBColor(245, 158, 11)      # Amber 500
    PURPLE = RGBColor(139, 92, 246)     # Violet 500
    
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    TEXT_LIGHT = RGBColor(241, 245, 249)
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CUSTOMER SEGMENTATION ML"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN
        p.font.name = "Segoe UI"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_BG)

    # Accent Top Line
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.5), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = CYAN
    accent.line.fill.background()

    # Title & Subtitle Box
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(3.0))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "MACHINE LEARNING PROJECT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.font.name = "Segoe UI"
    p0.space_after = Pt(12)

    p1 = tf.add_paragraph()
    p1.text = "Customer Segmentation using Unsupervised Learning"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Segoe UI"
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Advanced Behavioral Clustering, Dimensionality Reduction & Strategic Persona Profiling"
    p2.font.size = Pt(18)
    p2.font.color.rgb = TEXT_LIGHT_MUTED
    p2.font.name = "Segoe UI"

    # Presenter Card
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(5.5), Inches(1.6))
    card1.fill.solid()
    card1.fill.fore_color.rgb = DARK_CARD
    card1.line.color.rgb = RGBColor(51, 65, 85)

    tf_c = card1.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Author / Data Scientist"
    p.font.size = Pt(11)
    p.font.color.rgb = CYAN
    p.font.name = "Segoe UI"

    p = tf_c.add_paragraph()
    p.text = "Vandana Illipilla"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Segoe UI"

    p = tf_c.add_paragraph()
    p.text = "GitHub: @Vandana251  |  Domain: Retail & Customer Analytics"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    p.font.name = "Segoe UI"

    # Tech Stack Highlights Card
    card2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(5.0), Inches(5.7), Inches(1.6))
    card2.fill.solid()
    card2.fill.fore_color.rgb = DARK_CARD
    card2.line.color.rgb = RGBColor(51, 65, 85)

    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Core Tech Stack"
    p.font.size = Pt(11)
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"

    p = tf_c2.add_paragraph()
    p.text = "Python • Scikit-Learn • SciPy • Pandas • Seaborn"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Segoe UI"

    p = tf_c2.add_paragraph()
    p.text = "Algorithms: K-Means • Hierarchical (Ward) • DBSCAN • PCA"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 2: Project Overview & Objectives
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, LIGHT_BG)
    add_header(slide2, "Project Overview & Business Objectives")

    cards_data2 = [
        ("The Business Problem", "Traditional one-size-fits-all marketing leads to poor ROI, high customer churn, and irrelevant promotions. Understanding diverse customer behaviors is essential for targeted growth.", CORAL),
        ("The Machine Learning Solution", "Leverage unsupervised learning algorithms to automatically uncover hidden patterns, behavioral traits, and demographic groups without requiring manual labels.", TEAL),
        ("Expected Business Impact", "Personalized campaign targeting, optimized product cataloging, improved customer lifetime value (LTV), and data-driven loyalty programs.", PURPLE)
    ]

    for i, (title, desc, color) in enumerate(cards_data2):
        x = Inches(0.8 + i * 4.0)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        # Top indicator line
        top_line = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(0.12))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = color
        top_line.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 3: Dataset Architecture & Features
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, LIGHT_BG)
    add_header(slide3, "Dataset Architecture & Feature Matrix (10,000 Records)")

    feat_cards = [
        ("Demographic Profile", ["Birth Year & Customer Age", "Education (Graduation, PhD, etc.)", "Marital Status & Family Size", "Annual Income & Senior Citizen status"], CYAN),
        ("Spending Patterns", ["Wine & Luxury spending", "Meat, Fish & Fruit purchases", "Sweet & Confectionery spend", "Total Spend across categories"], CORAL),
        ("Channel Preferences", ["Web & Catalog purchases", "Physical Store visits", "Website activity / visits per month", "Discount & Deal usage frequency"], TEAL),
        ("Digital Engagement", ["Mobile App user status", "Email campaign subscriptions", "Device Type (Desktop / Mobile)", "Social Media usage & Referral flags"], AMBER)
    ]

    for i, (title, items, color) in enumerate(feat_cards):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.5 + row * 2.7)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_MUTED
            p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 4: Exploratory Data Analysis (EDA)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, LIGHT_BG)
    add_header(slide4, "Exploratory Data Analysis (EDA) Key Insights")

    eda_points = [
        ("Income vs Spend Correlation", "High-income brackets exhibit an exponential rise in luxury categories (Wines & Premium Meat), while middle brackets favor store discounts and regular staple products.", CYAN),
        ("Family Size Impact", "Customers with kids/teens demonstrate higher sensitivity to discounts, multiple web visits, and prioritize utility/bulk items over luxury goods.", CORAL),
        ("Digital vs Physical Channels", "Mobile app users have a 35% higher campaign acceptance rate compared to desktop-only shoppers, representing a high-potential digital audience.", TEAL),
        ("Customer Tenure & Loyalty", "Longer tenure customers maintain consistent store visits, but newer acquisitions demonstrate higher responsiveness to digital catalog promotions.", PURPLE)
    ]

    for i, (title, desc, color) in enumerate(eda_points):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.5 + row * 2.7)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 5: Data Preprocessing & Pipeline
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, LIGHT_BG)
    add_header(slide5, "Data Preprocessing & Feature Engineering Pipeline")

    steps5 = [
        ("1. Data Cleaning", "Handled missing data, converted dates (Customer_Since) to duration/tenure, and normalized categorical distributions.", CYAN),
        ("2. Feature Engineering", "Created aggregate metrics: Total_Amount_Spent, Total_Children, Customer_Age, and Channel_Diversity_Index.", TEAL),
        ("3. Outlier Management", "Applied IQR thresholding and robust filtering on income and extreme spend spikes to protect distance metrics.", AMBER),
        ("4. Feature Scaling", "Utilized StandardScaler to normalize numerical variables (mean=0, std=1) for Euclidean and density distances.", PURPLE)
    ]

    for i, (title, desc, color) in enumerate(steps5):
        x = Inches(0.8 + i * 2.95)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(2.75), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        top_line = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.6), Inches(2.75), Inches(0.12))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = color
        top_line.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 6: Dimensionality Reduction (PCA)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, LIGHT_BG)
    add_header(slide6, "Dimensionality Reduction via Principal Component Analysis (PCA)")

    # Left Column: Concepts
    card_l = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.75), Inches(5.2))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = CARD_BORDER

    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Why PCA is Essential for Segmentation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.font.name = "Segoe UI"
    p.space_after = Pt(12)

    points_l = [
        "Curse of Dimensionality: Distance metrics in high dimensions lose discriminative power.",
        "Multicollinearity Removal: Spending across categories and store visit counts correlate heavily.",
        "Noise Reduction: Focuses algorithm attention on primary variance-driving components.",
        "Intuitive 2D/3D Visualization: Allows clear boundary visualization for stakeholder reporting."
    ]
    for pt in points_l:
        p = tf_l.add_paragraph()
        p.text = f"•  {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    # Right Column: Variance Explained
    card_r = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.75), Inches(5.2))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_BG
    card_r.line.color.rgb = CARD_BORDER

    tf_r = card_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Principal Components Breakdown"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"
    p.space_after = Pt(12)

    components = [
        ("PC1: Purchasing Power & Luxury Spend", "Captures ~42% variance (Income, Wine, Meat spend)."),
        ("PC2: Channel & Tech Adoption", "Captures ~22% variance (Web, Mobile App vs Store)."),
        ("PC3: Family Dynamics & Age", "Captures ~14% variance (Children at home, Age profile)."),
        ("Cumulative Explained Variance", "> 75% total variance preserved in top components.")
    ]
    for title, desc in components:
        p = tf_r.add_paragraph()
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        
        p = tf_r.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    # ==========================================
    # SLIDE 7: Unsupervised Algorithms Comparison
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, LIGHT_BG)
    add_header(slide7, "Clustering Algorithms: Architecture & Comparison")

    models7 = [
        ("K-Means Clustering", "Centroid-Based", "Partitions dataset into K distinct spherical clusters by minimizing within-cluster sum of squares (WCSS). Fast, highly scalable, and intuitive.", CYAN),
        ("Agglomerative Clustering", "Hierarchical / Tree", "Bottom-up approach creating hierarchical trees (Dendrograms). Excellent for nested sub-groups and does not assume predefined shapes.", TEAL),
        ("DBSCAN", "Density-Based", "Groups points with dense neighborhoods and automatically isolates anomalies/outliers as noise (-1). Ideal for non-linear structures.", CORAL)
    ]

    for i, (name, tag, desc, color) in enumerate(models7):
        x = Inches(0.8 + i * 4.0)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        top_line = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(0.12))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = color
        top_line.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = name
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 8: Optimal Cluster Evaluation
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, LIGHT_BG)
    add_header(slide8, "Determining Optimal K: Validation Techniques")

    card_el = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.75), Inches(5.2))
    card_el.fill.solid()
    card_el.fill.fore_color.rgb = CARD_BG
    card_el.line.color.rgb = CARD_BORDER

    tf_el = card_el.text_frame
    tf_el.word_wrap = True
    p = tf_el.paragraphs[0]
    p.text = "1. Elbow Method & WCSS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.font.name = "Segoe UI"
    p.space_after = Pt(10)

    el_text = [
        "Evaluated K values ranging from 2 to 10.",
        "Identified sharp reduction in inertia followed by leveling off.",
        "Clear inflection point ('elbow') observed at K = 4.",
        "Provides foundational baseline for compact cluster partitioning."
    ]
    for pt in el_text:
        p = tf_el.add_paragraph()
        p.text = f"•  {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    card_sil = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.75), Inches(5.2))
    card_sil.fill.solid()
    card_sil.fill.fore_color.rgb = CARD_BG
    card_sil.line.color.rgb = CARD_BORDER

    tf_sil = card_sil.text_frame
    tf_sil.word_wrap = True
    p = tf_sil.paragraphs[0]
    p.text = "2. Silhouette Analysis & Dendrogram"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"
    p.space_after = Pt(10)

    sil_text = [
        "Silhouette coefficients validated cluster cohesion vs separation.",
        "K = 4 delivered balanced silhouette widths without negative misclassifications.",
        "Hierarchical Dendrogram cutting confirms 4 primary macro-branches.",
        "Optimal balance between statistical rigor and business interpretability."
    ]
    for pt in sil_text:
        p = tf_sil.add_paragraph()
        p.text = f"•  {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    # ==========================================
    # SLIDE 9: Customer Personas / Cluster Profiles
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, LIGHT_BG)
    add_header(slide9, "Discovered Customer Personas (4 Core Clusters)")

    personas = [
        ("Cluster 0: Elite Affluents", "High Income | High Spend", "Top spenders on premium Wine & Meat. Low discount dependency. High catalog & web channel usage.", CYAN),
        ("Cluster 1: Digital Young Pros", "Moderate Income | Tech Savvy", "Active mobile app users, responsive to email campaigns and social promotions. High web engagement.", TEAL),
        ("Cluster 2: Value Families", "Mid Income | Discount Focused", "Larger family size with children. High frequency of deals and store purchases. Price conscious.", AMBER),
        ("Cluster 3: Conservative Shoppers", "Low/Moderate | Traditional", "Infrequent digital activity. Preference for in-person physical store purchases with lower overall basket size.", PURPLE)
    ]

    for i, (p_name, p_tag, p_desc, color) in enumerate(personas):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.5 + row * 2.7)

        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        bar = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = p_name
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"

        p = tf.add_paragraph()
        p.text = p_tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = p_desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 10: Model Evaluation Metrics
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, LIGHT_BG)
    add_header(slide10, "Quantitative Model Evaluation & Performance")

    metrics_data = [
        ("Silhouette Score", "Higher is better (-1 to +1)", "Measures cohesion within clusters vs separation from neighboring clusters. K-Means achieved top score of 0.46.", CYAN),
        ("Calinski-Harabasz Index", "Higher is better", "Variance ratio criterion evaluating between-cluster dispersion to within-cluster dispersion.", TEAL),
        ("Davies-Bouldin Index", "Lower is better", "Calculates average similarity measure of each cluster with its most similar cluster.", CORAL)
    ]

    for i, (m_name, m_tag, m_desc, color) in enumerate(metrics_data):
        x = Inches(0.8 + i * 4.0)
        card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        top_line = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(0.12))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = color
        top_line.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = m_tag.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = m_name
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(12)

        p = tf.add_paragraph()
        p.text = m_desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 11: Business Applications & Strategy
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, LIGHT_BG)
    add_header(slide11, "Strategic Recommendations & Business Action Plan")

    strategies = [
        ("Elite Affluents: VIP Concierge", "Deploy exclusive private tastings, premium catalog access, high-tier loyalty perks, and dedicated personal account managers.", CYAN),
        ("Digital Pros: App Gamification", "Push notifications, flash sales, referral incentives, and personalized AI product recommendations on mobile.", TEAL),
        ("Value Families: Bundle Discounts", "Family bundle deals, loyalty cashbacks, seasonal school/holiday promotions, and multi-buy savings.", AMBER),
        ("Traditionalists: Store Engagement", "In-store experiential events, print coupons, localized flyers, and assisted checkout loyalty enrollment.", PURPLE)
    ]

    for i, (title, desc, color) in enumerate(strategies):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.5 + row * 2.7)

        card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 12: Conclusion & Future Roadmap (Dark Theme)
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, DARK_BG)

    # Accent Top Line
    accent12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(1.5), Inches(0.08))
    accent12.fill.solid()
    accent12.fill.fore_color.rgb = CYAN
    accent12.line.fill.background()

    t_box12 = slide12.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(1.0))
    tf12 = t_box12.text_frame
    tf12.word_wrap = True
    p = tf12.paragraphs[0]
    p.text = "Conclusion & Future Roadmap"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Segoe UI"

    card_c1 = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.5), Inches(5.75), Inches(4.2))
    card_c1.fill.solid()
    card_c1.fill.fore_color.rgb = DARK_CARD
    card_c1.line.color.rgb = RGBColor(51, 65, 85)

    tf = card_c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Project Takeaways"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.font.name = "Segoe UI"
    p.space_after = Pt(10)

    takeaways = [
        "Unsupervised ML successfully uncovered 4 distinct, highly actionable customer segments from 10,000 records.",
        "PCA significantly improved clustering separation and computational performance.",
        "Enables transition from generic marketing to hyper-personalized campaigns.",
        "Directly enhances customer lifetime value (LTV) and marketing ROI."
    ]
    for pt in takeaways:
        p = tf.add_paragraph()
        p.text = f"✔  {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    card_c2 = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.5), Inches(5.75), Inches(4.2))
    card_c2.fill.solid()
    card_c2.fill.fore_color.rgb = DARK_CARD
    card_c2.line.color.rgb = RGBColor(51, 65, 85)

    tf2 = card_c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Future Enhancements & Next Steps"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Segoe UI"
    p.space_after = Pt(10)

    future = [
        "Dynamic Stream Clustering: Incorporate real-time behavioral streams using MiniBatch K-Means.",
        "Recommendation Engine: Integrate collaborative filtering for product suggestions per cluster.",
        "Churn Prediction Integration: Build supervised classification models on top of cluster labels.",
        "Interactive Dashboard: Deploy Streamlit / PowerBI dashboard for executive monitoring."
    ]
    for pt in future:
        p = tf2.add_paragraph()
        p.text = f"🚀  {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    output_path = os.path.abspath("Customer_Segmentation_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
